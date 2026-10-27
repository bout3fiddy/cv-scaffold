#!/usr/bin/env python3
"""Build deterministic ATS-safe PDF and DOCX candidates from semantic HTML."""

from __future__ import annotations

import argparse
import html as html_module
import os
import re
import tempfile
import zipfile
from dataclasses import dataclass
from functools import partial
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import etree, html
from reportlab import rl_config
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfdoc, pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer


ROOT = Path(__file__).resolve().parent
CSS_PATH = ROOT / "resume.css"
FONT_DIR = ROOT / "fonts"
DEFAULT_OUTPUT_DIR = ROOT / "output"
FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
FIXED_DOCUMENT_TIME = "2026-09-10T00:00:00Z"
DOCX_FONT_NAME = "JetBrains Mono"

SOURCES = {
    "quant": {
        "html": ROOT / "examples/quantitative-research.html",
        "stem": "avery-quill-quant",
    },
    "swe": {
        "html": ROOT / "examples/software-engineering.html",
        "stem": "avery-quill-swe",
    },
    "strategy": {"html": ROOT / "examples/strategy.html", "stem": "avery-quill-strategy"},
}


def candidate_stem(key: str, revision: str = "v2") -> str:
    stem = SOURCES[key]["stem"]
    return stem if revision == "v2" else f"{stem}_V1"


@dataclass(frozen=True)
class ResumeSource:
    path: Path
    tree: html.HtmlElement
    article: html.HtmlElement
    variant: str
    title: str
    author: str
    subject: str
    expected_links: int


class ResumeCanvas(Canvas):
    def __init__(self, *args, title: str, author: str, subject: str, **kwargs) -> None:
        kwargs["invariant"] = 1
        kwargs["initialFontName"] = "JBMono"
        kwargs["initialFontSize"] = 8.25
        kwargs["initialLeading"] = 12.75
        kwargs["lang"] = "en-US"
        kwargs["cropBox"] = (0, 0, A4[0], A4[1])
        super().__init__(*args, **kwargs)
        self.setTitle(title)
        self.setAuthor(author)
        self.setSubject(subject)
        self.setCreator("CV starter semantic resume builder")
        self.setProducer("ReportLab deterministic ATS resume builder")
        self._doc.Catalog.Lang = pdfdoc.PDFString("en-US")


def parse_css_tokens(path: Path = CSS_PATH) -> dict[str, str]:
    text = path.read_text(encoding="utf-8")
    return {
        name: value.strip()
        for name, value in re.findall(r"--([a-z0-9-]+)\s*:\s*([^;]+);", text)
    }


def token_float(tokens: dict[str, str], name: str) -> float:
    return float(tokens[name].removesuffix("pt").strip())


def select(node: etree._Element, selector: str) -> list[etree._Element]:
    """Select the simple tag.class forms used by the resume sources."""
    tag, class_name = selector.split(".", 1)
    return node.xpath(
        f".//{tag}[contains(concat(' ', normalize-space(@class), ' '), ' {class_name} ')]"
    )


def load_source(path: Path) -> ResumeSource:
    tree = html.parse(str(path)).getroot()
    articles = select(tree, "article.resume")
    if len(articles) != 1:
        raise ValueError(f"Expected exactly one article.resume in {path}")
    article = articles[0]
    variant = article.get("data-variant")
    title_nodes = tree.xpath("//head/title")
    author_nodes = tree.xpath('//meta[@name="author"]/@content')
    subject_nodes = tree.xpath('//meta[@name="description"]/@content')
    return ResumeSource(
        path=path,
        tree=tree,
        article=article,
        variant=variant,
        title=title_nodes[0].text_content().strip(),
        author=author_nodes[0].strip(),
        subject=subject_nodes[0].strip(),
        expected_links=int(article.get("data-expected-links", "0")),
    )


def color_hex(tokens: dict[str, str], name: str) -> str:
    return tokens[name].lstrip("#").upper()


def reportlab_markup(node: etree._Element, tokens: dict[str, str]) -> str:
    pieces: list[str] = []

    def append_text(text: str | None, nowrap: bool = False) -> None:
        if not text:
            return
        escaped = html_module.escape(text)
        if nowrap:
            escaped = escaped.replace(" ", "&nbsp;")
        pieces.append(escaped)

    def walk(element: etree._Element) -> None:
        classes = set((element.get("class") or "").split())
        tag = element.tag.lower() if isinstance(element.tag, str) else ""
        nowrap = "nowrap" in classes or tag == "time"
        prefix = ""
        suffix = ""

        if tag in {"strong", "b"} or classes & {"role", "degree", "tech-label"}:
            prefix, suffix = '<font name="JBMono-Bold">', "</font>"
        elif tag in {"em", "i"} or classes & {"employer", "school", "date"}:
            prefix, suffix = '<font name="JBMono-Italic">', "</font>"

        tag_colors = {
            "python": "paper-python",
            "zig": "paper-zig",
            "typescript": "paper-typescript",
            "red": "paper-red",
            "purple": "paper-purple",
            "teal": "paper-teal",
            "gold": "paper-gold",
            "orange": "paper-orange",
        }
        for class_name, token_name in tag_colors.items():
            if class_name in classes:
                prefix += f'<font color="#{color_hex(tokens, token_name)}"><u>'
                suffix = "</u></font>" + suffix
                break

        if tag == "a":
            href = html_module.escape(element.get("href", ""), quote=True)
            prefix += f'<a href="{href}" color="#{color_hex(tokens, "paper-link")}"><u>'
            suffix = "</u></a>" + suffix

        pieces.append(prefix)
        append_text(element.text, nowrap=nowrap)
        for child in element:
            walk(child)
            append_text(child.tail)
        pieces.append(suffix)

    walk(node)
    return "".join(pieces).strip()


def source_text(node: etree._Element) -> str:
    return " ".join(node.text_content().split())


def register_fonts() -> None:
    fonts = {
        "JBMono": FONT_DIR / "JetBrainsMono-Regular.ttf",
        "JBMono-Bold": FONT_DIR / "JetBrainsMono-Bold.ttf",
        "JBMono-Italic": FONT_DIR / "JetBrainsMono-Italic.ttf",
    }
    missing = [str(path) for path in fonts.values() if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing required fonts: " + ", ".join(missing))
    for name, path in fonts.items():
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(path), asciiReadable=True))
    pdfmetrics.registerFontFamily(
        "JBMono",
        normal="JBMono",
        bold="JBMono-Bold",
        italic="JBMono-Italic",
        boldItalic="JBMono-Bold",
    )


def pdf_styles(tokens: dict[str, str], variant: str, revision: str = "v2") -> dict[str, ParagraphStyle]:
    body = token_float(tokens, "pdf-body-size")
    body_leading = token_float(tokens, "pdf-body-leading")
    if revision == "v2":
        body = token_float(tokens, "pdf-v2-body-size")
        body_leading = token_float(tokens, "pdf-v2-body-leading")
    elif variant == "swe":
        body = 8.1
        body_leading = 11.7
    section_size = token_float(tokens, "pdf-v2-section-size") if revision == "v2" else token_float(tokens, "pdf-section-size")
    section_leading = token_float(tokens, "pdf-v2-section-leading") if revision == "v2" else token_float(tokens, "pdf-section-leading")
    section_tracking = token_float(tokens, "pdf-v2-section-tracking") if revision == "v2" else 1.65
    common = {
        "alignment": TA_LEFT,
        "allowWidows": 0,
        "allowOrphans": 0,
        "splitLongWords": 0,
    }
    return {
        "title": ParagraphStyle(
            "ResumeTitle",
            **common,
            fontName="JBMono-Bold",
            fontSize=token_float(tokens, "pdf-title-size"),
            leading=token_float(tokens, "pdf-title-size"),
            textColor=colors.HexColor(tokens["paper-ink"]),
            spaceAfter=4.5,
        ),
        "contact": ParagraphStyle(
            "Contact",
            **common,
            fontName="JBMono",
            fontSize=token_float(tokens, "pdf-contact-size"),
            leading=10.5,
            textColor=colors.HexColor(tokens["paper-secondary"]),
        ),
        "summary": ParagraphStyle(
            "Summary",
            **common,
            fontName="JBMono",
            fontSize=body,
            leading=token_float(tokens, "pdf-v2-summary-leading") if revision == "v2" else (token_float(tokens, "pdf-summary-leading") if variant == "quant" else 12.8),
            textColor=colors.HexColor(tokens["paper-body"]),
        ),
        "section": ParagraphStyle(
            "Section",
            **common,
            fontName="JBMono-Bold",
            fontSize=section_size,
            leading=section_leading,
            textColor=colors.HexColor(tokens["paper-muted"]),
            tracking=section_tracking,
            keepWithNext=True,
        ),
        "body": ParagraphStyle(
            "Body",
            **common,
            fontName="JBMono",
            fontSize=body,
            leading=body_leading,
            textColor=colors.HexColor(tokens["paper-body"]),
        ),
        "job_heading": ParagraphStyle(
            "JobHeading",
            **common,
            fontName="JBMono",
            fontSize=token_float(tokens, "pdf-role-size"),
            leading=token_float(tokens, "pdf-role-leading"),
            textColor=colors.HexColor(tokens["paper-secondary"]),
            keepWithNext=True,
        ),
        "tech": ParagraphStyle(
            "Tech",
            **common,
            fontName="JBMono",
            fontSize=token_float(tokens, "pdf-tech-size"),
            leading=token_float(tokens, "pdf-tech-leading"),
            textColor=colors.HexColor(tokens["paper-secondary"]),
            keepWithNext=True,
        ),
        "bullet": ParagraphStyle(
            "Bullet",
            **common,
            fontName="JBMono",
            fontSize=body,
            leading=body_leading,
            textColor=colors.HexColor(tokens["paper-body"]),
            leftIndent=12,
            firstLineIndent=-12,
            spaceBefore=2.0,
        ),
        "compact": ParagraphStyle(
            "Compact",
            **common,
            fontName="JBMono",
            fontSize=body,
            leading=body_leading,
            textColor=colors.HexColor(tokens["paper-body"]),
            spaceBefore=2.0,
        ),
    }


def build_pdf(source: ResumeSource, output_path: Path, tokens: dict[str, str], revision: str = "v2") -> None:
    register_fonts()
    styles = pdf_styles(tokens, source.variant, revision)
    margin_name = "pdf-quant-margin" if source.variant == "quant" else "pdf-swe-margin"
    margin = token_float(tokens, margin_name)
    vertical_margin = token_float(tokens, "pdf-v2-swe-vertical-margin") if revision == "v2" and source.variant == "swe" else margin
    # SimpleDocTemplate adds 6 pt of frame padding on every edge. Compensate so
    # the visible content starts at the Paper artboard's 30/36 pt inset.
    flow_margin = margin - 6
    flow_vertical_margin = vertical_margin - 6
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        leftMargin=flow_margin,
        rightMargin=flow_margin,
        topMargin=flow_vertical_margin,
        bottomMargin=flow_vertical_margin,
        title=source.title,
        author=source.author,
        subject=source.subject,
        invariant=1,
        pageCompression=1,
    )
    story: list = []
    h1 = source.article.xpath(".//h1")[0]
    contact = select(source.article, "p.contact")[0]
    summary = select(source.article, "p.summary")[0]
    story.append(Paragraph(source_text(h1), styles["title"]))
    story.append(Paragraph(reportlab_markup(contact, tokens), styles["contact"]))
    header_summary_gap = token_float(tokens, "pdf-v2-header-summary-gap") if revision == "v2" else (13.5 if source.variant == "quant" else 16.5)
    story.append(Spacer(1, header_summary_gap))
    story.append(Paragraph(reportlab_markup(summary, tokens), styles["summary"]))

    for section in source.article.xpath("./section"):
        heading = section.xpath("./h2")[0]
        section_class = set((section.get("class") or "").split())
        if revision == "v2":
            before = token_float(tokens, "pdf-v2-section-gap")
        elif section_class & {"selected-work"}:
            before = 13.5
        elif source_text(heading).lower() in {"education", "publications"}:
            before = 10.5
        else:
            before = 8.0 if source.variant == "quant" else 10.0
        story.append(Spacer(1, before))
        story.append(Paragraph(source_text(heading).upper(), styles["section"]))
        story.append(Spacer(1, token_float(tokens, "pdf-v2-section-content-gap") if revision == "v2" else 4.0))

        if source_text(heading).lower() == "professional experience":
            jobs = section.xpath("./article[contains(concat(' ', normalize-space(@class), ' '), ' job ')]")
            for index, job in enumerate(jobs):
                job_parts: list = []
                if index:
                    gap_name = "pdf-v2-job-gap" if revision == "v2" else None
                    job_gap = token_float(tokens, gap_name) if gap_name else (6.0 if source.variant == "quant" else 5.5)
                    job_parts.append(Spacer(1, job_gap))
                head = select(job, "p.job-heading")[0]
                job_parts.append(Paragraph(reportlab_markup(head, tokens), styles["job_heading"]))
                if revision == "v2":
                    job_parts.append(Spacer(1, token_float(tokens, "pdf-v2-role-content-gap")))
                else:
                    tech = select(job, "p.tech")[0]
                    job_parts.append(Spacer(1, token_float(tokens, "pdf-role-tech-gap")))
                    job_parts.append(Paragraph(reportlab_markup(tech, tokens), styles["tech"]))
                    job_parts.append(Spacer(1, token_float(tokens, "pdf-tech-content-gap")))
                for bullet in select(job, "p.bullet"):
                    job_parts.append(Paragraph(reportlab_markup(bullet, tokens), styles["bullet"]))
                story.append(KeepTogether(job_parts))
        else:
            for paragraph in section.xpath("./p"):
                style = styles["body"] if "skills-copy" in (paragraph.get("class") or "") else styles["compact"]
                story.append(Paragraph(reportlab_markup(paragraph, tokens), style))

    canvasmaker = partial(
        ResumeCanvas,
        title=source.title,
        author=source.author,
        subject=source.subject,
    )
    doc.build(story, canvasmaker=canvasmaker)


def set_run_font(run, name: str, size: float, color: str, bold: bool = False, italic: bool = False, underline: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.bold = bold
    run.font.italic = italic
    run.font.underline = underline
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)
    language = OxmlElement("w:lang")
    language.set(qn("w:val"), "en-US")
    run._element.get_or_add_rPr().append(language)


def append_hyperlink(paragraph, text: str, url: str, size: float, font_name: str, color: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), font_name)
    properties.append(fonts)
    size_element = OxmlElement("w:sz")
    size_element.set(qn("w:val"), str(round(size * 2)))
    properties.append(size_element)
    color_element = OxmlElement("w:color")
    color_element.set(qn("w:val"), color)
    properties.append(color_element)
    underline_element = OxmlElement("w:u")
    underline_element.set(qn("w:val"), "single")
    properties.append(underline_element)
    language = OxmlElement("w:lang")
    language.set(qn("w:val"), "en-US")
    properties.append(language)
    run.append(properties)
    text_element = OxmlElement("w:t")
    if text[:1].isspace() or text[-1:].isspace():
        text_element.set(qn("xml:space"), "preserve")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def append_docx_node(paragraph, node: etree._Element, tokens: dict[str, str], size: float) -> None:
    default_color = color_hex(tokens, "paper-body")

    def append_text(text: str | None, settings: dict[str, object]) -> None:
        # Pretty-printed HTML contributes indentation-only text nodes between
        # inline elements. Skip these formatting-only nodes.
        if not text or (not text.strip() and ("\n" in text or "\r" in text)):
            return
        run = paragraph.add_run(text)
        set_run_font(
            run,
            DOCX_FONT_NAME,
            size,
            str(settings.get("color", default_color)),
            bold=bool(settings.get("bold", False)),
            italic=bool(settings.get("italic", False)),
            underline=bool(settings.get("underline", False)),
        )

    def walk(element: etree._Element, inherited: dict[str, object]) -> None:
        classes = set((element.get("class") or "").split())
        tag = element.tag.lower() if isinstance(element.tag, str) else ""
        settings = dict(inherited)
        if tag in {"strong", "b"} or classes & {"role", "degree", "tech-label"}:
            settings["bold"] = True
        if tag in {"em", "i"} or classes & {"employer", "school", "date"}:
            settings["italic"] = True
        if "date" in classes:
            settings["color"] = color_hex(tokens, "paper-muted")
        tag_colors = {
            "python": "paper-python",
            "zig": "paper-zig",
            "typescript": "paper-typescript",
            "red": "paper-red",
            "purple": "paper-purple",
            "teal": "paper-teal",
            "gold": "paper-gold",
            "orange": "paper-orange",
        }
        for class_name, token_name in tag_colors.items():
            if class_name in classes:
                settings["color"] = color_hex(tokens, token_name)
                settings["underline"] = True
                break

        if tag == "a":
            if element.text:
                append_hyperlink(
                    paragraph,
                    element.text,
                    element.get("href", ""),
                    size,
                    DOCX_FONT_NAME,
                    color_hex(tokens, "paper-link"),
                )
            for child in element:
                walk(child, settings)
                append_text(child.tail, inherited)
            return

        append_text(element.text, settings)
        for child in element:
            walk(child, settings)
            append_text(child.tail, settings)

    walk(node, {})


def configure_paragraph(paragraph, size: float, leading: float, before: float = 0, after: float = 0, keep_next: bool = False) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = Pt(leading)
    fmt.keep_with_next = keep_next
    fmt.widow_control = True
    paragraph.style = "Resume Body"


def add_docx_paragraph(document: Document, node: etree._Element, tokens: dict[str, str], size: float, leading: float, before: float = 0, after: float = 0, keep_next: bool = False, left_indent: float = 0, first_line_indent: float = 0):
    paragraph = document.add_paragraph()
    configure_paragraph(paragraph, size, leading, before, after, keep_next)
    if left_indent:
        paragraph.paragraph_format.left_indent = Pt(left_indent)
    if first_line_indent:
        paragraph.paragraph_format.first_line_indent = Pt(first_line_indent)
    append_docx_node(paragraph, node, tokens, size)
    return paragraph


def normalize_docx(raw_path: Path, output_path: Path) -> None:
    with zipfile.ZipFile(raw_path, "r") as source_zip:
        entries = {name: source_zip.read(name) for name in source_zip.namelist()}
    core_name = "docProps/core.xml"
    if core_name in entries:
        core = entries[core_name].decode("utf-8")
        core = re.sub(
            r"<dcterms:(created|modified)[^>]*>.*?</dcterms:\1>",
            lambda match: re.sub(r">.*?<", f">{FIXED_DOCUMENT_TIME}<", match.group(0)),
            core,
        )
        entries[core_name] = core.encode("utf-8")
    with zipfile.ZipFile(output_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target_zip:
        for name in sorted(entries):
            info = zipfile.ZipInfo(name, date_time=FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target_zip.writestr(info, entries[name])


def build_docx(source: ResumeSource, output_path: Path, tokens: dict[str, str], revision: str = "v2") -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    margin = token_float(tokens, "pdf-quant-margin" if source.variant == "quant" else "pdf-swe-margin")
    vertical_margin = token_float(tokens, "pdf-v2-swe-vertical-margin") if revision == "v2" and source.variant == "swe" else margin
    section.top_margin = Pt(vertical_margin)
    section.bottom_margin = Pt(vertical_margin)
    section.left_margin = Pt(margin)
    section.right_margin = Pt(margin)
    section.start_type = WD_SECTION.NEW_PAGE

    styles = document.styles
    if "Resume Body" not in styles:
        resume_style = styles.add_style("Resume Body", WD_STYLE_TYPE.PARAGRAPH)
    else:
        resume_style = styles["Resume Body"]
    resume_style.base_style = styles["Normal"]
    resume_style.font.name = DOCX_FONT_NAME
    resume_style.font.size = Pt(token_float(tokens, "pdf-body-size"))
    resume_style.font.color.rgb = RGBColor.from_string(color_hex(tokens, "paper-body"))
    resume_style._element.rPr.rFonts.set(qn("w:ascii"), DOCX_FONT_NAME)
    resume_style._element.rPr.rFonts.set(qn("w:hAnsi"), DOCX_FONT_NAME)
    resume_style._element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT_NAME)
    resume_style.paragraph_format.left_indent = Pt(0)
    resume_style.paragraph_format.right_indent = Pt(0)
    resume_style.paragraph_format.first_line_indent = Pt(0)
    resume_style.paragraph_format.space_before = Pt(0)
    resume_style.paragraph_format.space_after = Pt(0)

    title_style = styles["Title"]
    title_style.base_style = styles["Normal"]
    title_style.font.name = DOCX_FONT_NAME
    title_style.font.size = Pt(token_float(tokens, "pdf-title-size"))
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string(color_hex(tokens, "paper-ink"))
    title_style._element.rPr.rFonts.set(qn("w:ascii"), DOCX_FONT_NAME)
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), DOCX_FONT_NAME)
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT_NAME)
    title_style.paragraph_format.left_indent = Pt(0)
    title_style.paragraph_format.right_indent = Pt(0)
    title_style.paragraph_format.first_line_indent = Pt(0)
    title_style.paragraph_format.space_before = Pt(0)
    title_style.paragraph_format.space_after = Pt(0)
    title_properties = title_style._element.get_or_add_pPr()
    for border in title_properties.findall(qn("w:pBdr")):
        title_properties.remove(border)

    properties = document.core_properties
    properties.title = source.title
    properties.author = source.author
    properties.subject = source.subject
    properties.keywords = "resume, ATS, quantitative research" if source.variant == "quant" else "resume, ATS, software engineering"
    properties.comments = "Built deterministically from semantic HTML derived from Paper design source."

    body_size = token_float(tokens, "pdf-v2-body-size") if revision == "v2" else (token_float(tokens, "pdf-body-size") if source.variant == "quant" else 8.1)
    body_leading = token_float(tokens, "pdf-v2-body-leading") if revision == "v2" else (token_float(tokens, "pdf-body-leading") if source.variant == "quant" else 11.7)
    h1_node = source.article.xpath(".//h1")[0]
    title_paragraph = document.add_paragraph()
    title_paragraph.style = "Title"
    title_paragraph.paragraph_format.space_after = Pt(4.5)
    title_paragraph.paragraph_format.line_spacing = Pt(token_float(tokens, "pdf-title-size"))
    title_run = title_paragraph.add_run(source_text(h1_node))
    set_run_font(
        title_run,
        DOCX_FONT_NAME,
        token_float(tokens, "pdf-title-size"),
        color_hex(tokens, "paper-ink"),
        bold=True,
    )
    add_docx_paragraph(
        document,
        select(source.article, "p.contact")[0],
        tokens,
        token_float(tokens, "pdf-contact-size"),
        10.5,
    )
    add_docx_paragraph(
        document,
        select(source.article, "p.summary")[0],
        tokens,
        body_size,
        token_float(tokens, "pdf-v2-summary-leading") if revision == "v2" else (token_float(tokens, "pdf-summary-leading") if source.variant == "quant" else 12.8),
        before=token_float(tokens, "pdf-v2-header-summary-gap") if revision == "v2" else (13.5 if source.variant == "quant" else 16.5),
    )

    for section_node in source.article.xpath("./section"):
        heading = section_node.xpath("./h2")[0]
        heading_text = source_text(heading)
        classes = set((section_node.get("class") or "").split())
        if revision == "v2":
            before = token_float(tokens, "pdf-v2-section-gap")
        elif "selected-work" in classes:
            before = 13.5
        elif heading_text.lower() in {"education", "publications"}:
            before = 10.5
        else:
            before = 8.0 if source.variant == "quant" else 10.0
        paragraph = document.add_paragraph()
        paragraph.style = "Resume Body"
        paragraph.paragraph_format.space_before = Pt(before)
        paragraph.paragraph_format.space_after = Pt(token_float(tokens, "pdf-v2-section-content-gap") if revision == "v2" else 4)
        paragraph.paragraph_format.line_spacing = Pt(token_float(tokens, "pdf-v2-section-leading") if revision == "v2" else token_float(tokens, "pdf-section-leading"))
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(heading_text.upper())
        set_run_font(
            run,
            "JetBrains Mono",
            token_float(tokens, "pdf-v2-section-size") if revision == "v2" else token_float(tokens, "pdf-section-size"),
            color_hex(tokens, "paper-muted"),
            bold=True,
        )
        run._element.get_or_add_rPr().append(OxmlElement("w:spacing"))
        run._element.rPr[-1].set(qn("w:val"), "18" if revision == "v2" else "26")

        if heading_text.lower() == "professional experience":
            jobs = section_node.xpath("./article[contains(concat(' ', normalize-space(@class), ' '), ' job ')]")
            for index, job in enumerate(jobs):
                add_docx_paragraph(
                    document,
                    select(job, "p.job-heading")[0],
                    tokens,
                    token_float(tokens, "pdf-role-size"),
                    token_float(tokens, "pdf-role-leading"),
                    before=(token_float(tokens, "pdf-v2-job-gap") if revision == "v2" else 6.0) if index else 0,
                    after=token_float(tokens, "pdf-v2-role-content-gap") if revision == "v2" else token_float(tokens, "pdf-role-tech-gap"),
                    keep_next=True,
                )
                if revision != "v2":
                    add_docx_paragraph(
                        document,
                        select(job, "p.tech")[0],
                        tokens,
                        token_float(tokens, "pdf-tech-size"),
                        token_float(tokens, "pdf-tech-leading"),
                        after=token_float(tokens, "pdf-tech-content-gap"),
                        keep_next=True,
                    )
                for bullet in select(job, "p.bullet"):
                    add_docx_paragraph(
                        document,
                        bullet,
                        tokens,
                        body_size,
                        body_leading,
                        before=2,
                        left_indent=12,
                        first_line_indent=-12,
                    )
        else:
            for paragraph_node in section_node.xpath("./p"):
                add_docx_paragraph(
                    document,
                    paragraph_node,
                    tokens,
                    body_size,
                    body_leading,
                    before=2 if heading_text.lower() != "skills" else 0,
                )

    settings = document.settings._element
    language = OxmlElement("w:themeFontLang")
    language.set(qn("w:val"), "en-US")
    settings.append(language)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="resume-docx-") as temp_dir:
        raw_path = Path(temp_dir) / "raw.docx"
        document.save(raw_path)
        normalize_docx(raw_path, output_path)


def build_variant(key: str, output_dir: Path, formats: Iterable[str], revision: str = "v2") -> list[Path]:
    spec = SOURCES[key]
    source = load_source(spec["html"])
    tokens = parse_css_tokens()
    if key in {"swe", "strategy"}:
        if revision != "v2":
            raise ValueError("The current standard and strategy examples require revision v2.")
        import swe_layout
        tokens = swe_layout.layout_tokens(tokens)
    outputs: list[Path] = []
    stem = candidate_stem(key, revision)
    if "pdf" in formats:
        pdf_path = output_dir / f"{stem}.pdf"
        if key in {"swe", "strategy"}:
            if key == "strategy":
                import strategy_layout
                strategy_layout.build_pdf(source, pdf_path, tokens)
            else:
                swe_layout.build_pdf(source, pdf_path, tokens)
        else:
            build_pdf(source, pdf_path, tokens, revision)
        outputs.append(pdf_path)
    if "docx" in formats:
        docx_path = output_dir / f"{stem}.docx"
        if key in {"swe", "strategy"}:
            swe_layout.build_docx(source, docx_path, tokens)
            if key == "strategy":
                import strategy_layout
                document = Document(docx_path)
                strategy_layout.align_docx_dates(document, source, tokens)
                with tempfile.TemporaryDirectory() as folder:
                    raw = Path(folder) / "raw.docx"
                    document.save(raw)
                    normalize_docx(raw, docx_path)
        else:
            build_docx(source, docx_path, tokens, revision)
        outputs.append(docx_path)
    return outputs


def main() -> None:
    parser = argparse.ArgumentParser()
    group = parser.add_mutually_exclusive_group()
    group.add_argument("--all", action="store_true")
    group.add_argument("--variant", choices=sorted(SOURCES))
    parser.add_argument("--format", choices=["pdf", "docx", "both"], default="both")
    parser.add_argument("--revision", choices=["v2"], default="v2")
    parser.add_argument("--source", type=Path, help="Custom HTML source for the selected variant")
    parser.add_argument("--stem", help="Output filename without extension")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    args = parser.parse_args()

    os.environ.setdefault("SOURCE_DATE_EPOCH", "1788998400")
    rl_config.invariant = 1
    formats = ("pdf", "docx") if args.format == "both" else (args.format,)
    variants = sorted(SOURCES) if args.all else [args.variant or "swe"]
    if args.all and (args.source or args.stem):
        parser.error("--source and --stem require a single variant")
    if args.source:
        SOURCES[variants[0]]["html"] = args.source
    if args.stem:
        SOURCES[variants[0]]["stem"] = args.stem
    args.output_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    for variant in variants:
        written.extend(build_variant(variant, args.output_dir, formats, args.revision))
    for path in written:
        print(path)


if __name__ == "__main__":
    main()
