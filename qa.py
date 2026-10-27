"""Portable source-to-export checks; no candidate-specific assertions."""
import argparse
import json
import re
from pathlib import Path

from docx import Document
import pdfplumber
from pypdf import PdfReader

import build


def normalized(text):
    return re.sub(r'\s+', '', text).casefold()


def check_pdf(path, paragraphs, expected_pages, urls=()):
    reader = PdfReader(path)
    assert len(reader.pages) == expected_pages, f'{path.name}: expected {expected_pages} pages'
    text = normalized(''.join(page.extract_text() for page in reader.pages))
    for paragraph in paragraphs:
        assert normalized(paragraph) in text, f'{path.name}: missing content: {paragraph[:90]}'
    for page in reader.pages:
        assert abs(float(page.mediabox.width) - build.A4[0]) < 1
        assert abs(float(page.mediabox.height) - build.A4[1]) < 1
        for font in page['/Resources']['/Font'].get_object().values():
            font = font.get_object()
            if '/FontDescriptor' in font:
                descriptor = font['/FontDescriptor']
                assert any(key in descriptor for key in ('/FontFile', '/FontFile2', '/FontFile3'))
    actual_urls = [str(a.get_object().get('/A', {}).get('/URI', ''))
                   for page in reader.pages for a in page.get('/Annots', [])]
    for url in urls:
        assert url in actual_urls, f'{path.name}: missing link {url}'
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for char in page.chars:
                assert char['x0'] >= 28 and char['x1'] <= page.width - 28, f'{path.name}: horizontal overflow'
                assert char['top'] >= 28 and char['bottom'] <= page.height - 28, f'{path.name}: vertical overflow'
    return {'file': path.name, 'pages': len(reader.pages), 'paragraphs_checked': len(paragraphs)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', type=Path, default=build.DEFAULT_OUTPUT_DIR)
    parser.add_argument('--pages', type=int, default=1)
    parser.add_argument('--personal', action='store_true')
    parser.add_argument('--source', type=Path, help='Custom HTML CV or plain-text letter')
    parser.add_argument('--pdf', type=Path, help='PDF for the custom source')
    parser.add_argument('--docx', type=Path, help='Optional Word export for a custom CV')
    parser.add_argument('--kind', choices=['cv', 'letter'], default='cv')
    args = parser.parse_args()
    if bool(args.source) != bool(args.pdf):
        parser.error('--source and --pdf must be provided together')
    if args.docx and (not args.source or args.kind != 'cv'):
        parser.error('--docx requires a custom CV source')
    if args.source:
        if args.kind == 'cv':
            source = build.load_source(args.source)
            paragraphs = [build.source_text(node) for node in source.article.xpath('.//h1 | .//h2 | .//p')]
            urls = source.article.xpath('.//a/@href')
        else:
            paragraphs = args.source.read_text().strip().split('\n\n')
            urls = []
        report = check_pdf(args.pdf, paragraphs, args.pages, urls)
        if args.docx:
            text = normalized(''.join(p.text for p in Document(args.docx).paragraphs))
            for paragraph in paragraphs:
                assert normalized(paragraph) in text, f'Word export missing: {paragraph[:90]}'
        if args.personal:
            assert (build.ROOT / 'profile/ground-truth.md').exists(), 'Complete profile/ground-truth.md first'
            content = args.source.read_text() + str(PdfReader(args.pdf).metadata)
            assert not re.search(r'Avery Quill|Lantern Studio|Mosslight Tools|Kitebox Workshop|Westmere College|Example Learning Institute|Paperboat Works|example\.com|fictional|\[[^\]]+\]', content, re.I), 'Replace example details and metadata'
        print(json.dumps(report, indent=2))
        print('PASS. Visual review and factual verification remain manual.')
        return
    reports = []
    if args.personal:
        assert (build.ROOT / 'profile/ground-truth.md').exists(), 'Complete profile/ground-truth.md first'
    for variant, spec in build.SOURCES.items():
        source = build.load_source(spec['html'])
        nodes = source.article.xpath('.//h1 | .//h2 | .//p')
        paragraphs = [build.source_text(node) for node in nodes]
        urls = source.article.xpath('.//a/@href')
        reports.append(check_pdf(args.output_dir / (spec['stem'] + '.pdf'), paragraphs, args.pages, urls))
        docx = Document(args.output_dir / (spec['stem'] + '.docx'))
        text = normalized(''.join(p.text for p in docx.paragraphs))
        for paragraph in paragraphs:
            assert normalized(paragraph) in text, f'{variant} DOCX missing: {paragraph[:90]}'
        if args.personal:
            contents = source.path.read_text()
            assert not re.search(r'Avery Quill|Lantern Studio|Mosslight Tools|Kitebox Workshop|Westmere College|Example Learning Institute|example\.com|fictional|\[[^\]]+\]', contents, re.I), f'{variant}: replace fictional details'
    letter = build.ROOT / 'examples/cover-letter.md'
    reports.append(check_pdf(args.output_dir / 'avery-quill-cover-letter.pdf', letter.read_text().strip().split('\n\n'), 1))
    if args.personal:
        assert not re.search(r'Avery Quill|example\.com|fictional', letter.read_text(), re.I), 'Replace fictional letter'
    (args.output_dir / 'qa.json').write_text(json.dumps(reports, indent=2) + '\n')
    print(json.dumps(reports, indent=2))
    print('PASS. Visual review and factual verification remain manual.')


if __name__ == '__main__':
    main()
