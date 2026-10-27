"""Approved two-page default resume layout, independent of any role folder."""
from pathlib import Path
from functools import partial
from html import escape
import tempfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.colors import HexColor
from reportlab.pdfbase.pdfmetrics import stringWidth, getFont
from struct import unpack_from
from reportlab.platypus import Flowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer

import build

STYLE_PATH = Path(__file__).with_name("swe.css")


def measure(tokens, name):
    return build.token_float(tokens, "swe-" + name)


def append_docx_markup(paragraph, node, tokens, size, inherited=None):
    """Preserve explicit emphasis and blue, non-underlined hyperlinks."""
    settings = dict(inherited or {})
    tag = node.tag if isinstance(node.tag, str) else ""
    if tag in {"strong", "b"}:
        settings["bold"] = True
    if tag in {"em", "i", "time"}:
        settings["italic"] = True
    if tag == "u":
        settings["underline"] = True
    if "zero-to-one" in node.get("class", "").split():
        settings["color"] = build.color_hex(tokens, "swe-zero-to-one")

    def add_text(text, formatting):
        if text:
            run = paragraph.add_run(text)
            build.set_run_font(
                run, "JetBrains Mono", size,
                formatting.get("color", build.color_hex(tokens, "paper-body")),
                bold=formatting.get("bold", False),
                italic=formatting.get("italic", False),
                underline=formatting.get("underline", False),
            )

    if tag == "a":
        build.append_hyperlink(paragraph, node.text_content(), node.get("href"),
                               size, "JetBrains Mono", build.color_hex(tokens, "paper-link"))
        properties = paragraph._p[-1].find(qn("w:r")).find(qn("w:rPr"))
        properties.find(qn("w:u")).set(qn("w:val"), "none")
        for flag in ("bold", "italic"):
            if settings.get(flag):
                properties.append(OxmlElement("w:b" if flag == "bold" else "w:i"))
        return
    if "skill-tag" in node.get("class", "").split():
        add_text(" " + node.text_content() + " ", {**settings, "color": build.color_hex(tokens, "swe-tag-text")})
        border = OxmlElement("w:bdr")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "2")
        border.set(qn("w:color"), build.color_hex(tokens, "swe-tag-border"))
        paragraph.runs[-1]._element.get_or_add_rPr().append(border)
        return
    add_text(node.text, settings)
    for child in node:
        append_docx_markup(paragraph, child, tokens, size, settings)
        add_text(child.tail, settings)


def build_docx(source, output_path, tokens):
    """Editable fallback using the same content and type scale as the PDF."""
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Pt(build.A4[0]), Pt(build.A4[1])
    section.left_margin = section.right_margin = Pt(measure(tokens, "horizontal-margin"))
    section.top_margin = section.bottom_margin = Pt(measure(tokens, "vertical-margin"))
    normal = document.styles["Normal"]
    normal.font.name = "JetBrains Mono"
    normal.font.size = Pt(measure(tokens, "body-size"))
    normal.paragraph_format.space_after = Pt(0)

    def paragraph(node, size=None, leading=None, before=0, after=0,
                  left=0, first=0, keep_next=False, text=None, bold=False, color=None):
        size = measure(tokens, "body-size") if size is None else size
        leading = measure(tokens, "body-leading") if leading is None else leading
        p = document.add_paragraph()
        fmt = p.paragraph_format
        fmt.line_spacing = Pt(leading)
        fmt.space_before, fmt.space_after = Pt(before), Pt(after)
        fmt.left_indent, fmt.first_line_indent = Pt(left), Pt(first)
        fmt.keep_with_next = keep_next
        fmt.keep_together = True
        fmt.widow_control = True
        if text is not None:
            build.set_run_font(p.add_run(text), "JetBrains Mono", size,
                               color or build.color_hex(tokens, "paper-body"), bold=bold)
        else:
            append_docx_markup(p, node, tokens, size, {"color": color} if color else None)
        return p

    paragraph(None, size=measure(tokens, "title-size"), leading=measure(tokens, "title-size"),
              after=measure(tokens, "title-gap"), keep_next=True,
              text=build.source_text(source.article.xpath(".//h1")[0]), bold=True,
              color=build.color_hex(tokens, "paper-ink"))
    for index, contact in enumerate(build.select(source.article, "p.contact")):
        paragraph(contact, size=measure(tokens, "contact-size"), leading=measure(tokens, "contact-leading"),
                  before=measure(tokens, "contact-gap") if index else 0, keep_next=True,
                  color=build.color_hex(tokens, "paper-secondary"))
    paragraph(build.select(source.article, "p.summary")[0], before=measure(tokens, "summary-gap"))
    for section_node in source.article.xpath("./section"):
        heading = section_node.xpath("./h2")[0]
        paragraph(None, size=measure(tokens, "section-size"), leading=measure(tokens, "section-leading"),
                  before=measure(tokens, "section-gap"), after=measure(tokens, "section-content-gap"), keep_next=True,
                  text=build.source_text(heading).upper(), bold=True,
                  color=build.color_hex(tokens, "paper-muted"))
        jobs = section_node.xpath("./article[@class='job']")
        if not jobs:
            for index, node in enumerate(section_node.xpath("./p")):
                paragraph(node, before=measure(tokens, "bullet-gap")
                          if index and "education-entry" in node.get("class", "") else 0)
        for index, job in enumerate(jobs):
            paragraph(build.select(job, "p.job-heading")[0], before=measure(tokens, "job-gap") if index else 0,
                      after=measure(tokens, "role-content-gap"), keep_next=True)
            for bullet_index, bullet in enumerate(build.select(job, "p.bullet")):
                classes = bullet.get("class", "").split()
                paragraph(bullet, before=measure(tokens, "bullet-gap") if bullet_index else 0,
                          left=measure(tokens, "sub-bullet-indent" if "sub-bullet" in classes else "bullet-indent"),
                          first=-measure(tokens, "bullet-indent"), keep_next="subheading" in classes)
    document.core_properties.title = source.title
    document.core_properties.author = source.author
    document.core_properties.subject = source.subject
    document.core_properties.keywords = "resume, software engineering, applied AI"
    document.core_properties.comments = "Built from the canonical semantic resume source."
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="resume-swe-docx-") as directory:
        raw = Path(directory) / "resume.docx"
        document.save(raw)
        build.normalize_docx(raw, output_path)


def layout_tokens(tokens):
    tokens = dict(tokens)
    tokens.update(build.parse_css_tokens(STYLE_PATH))
    mapping = {
        "pdf-title-size": "title-size",
        "pdf-contact-size": "contact-size",
        "pdf-v2-body-size": "body-size",
        "pdf-v2-body-leading": "body-leading",
        "pdf-v2-summary-leading": "body-leading",
        "pdf-v2-section-size": "section-size",
        "pdf-v2-section-leading": "section-leading",
        "pdf-v2-swe-vertical-margin": "vertical-margin",
        "pdf-role-size": "body-size",
        "pdf-role-leading": "body-leading",
        "pdf-v2-header-summary-gap": "summary-gap",
        "pdf-v2-section-gap": "section-gap",
        "pdf-v2-section-content-gap": "section-content-gap",
        "pdf-v2-job-gap": "job-gap",
        "pdf-v2-role-content-gap": "role-content-gap",
    }
    tokens.update({target: tokens["swe-" + source] for target, source in mapping.items()})
    return tokens


def markup(node, tokens):
    """Render explicit emphasis from the canonical source without implicit bold."""
    if not isinstance(node.tag, str):
        return ""
    content = escape(node.text or "")
    for child in node:
        content += markup(child, tokens) + escape(child.tail or "")
    if node.tag == "time":
        content = content.replace(" ", "&nbsp;")
    if node.tag in {"strong", "b"}:
        return f'<font name="JBMono-Bold">{content}</font>'
    if node.tag in {"em", "i", "time"}:
        return f'<font name="JBMono-Italic">{content}</font>'
    if node.tag == "u":
        return f"<u>{content}</u>"
    if node.tag == "span" and "zero-to-one" in node.get("class", "").split():
        return f'<font color="{tokens["swe-zero-to-one"]}">{content}</font>'
    if node.tag == "a":
        href = escape(node.get("href", ""), quote=True)
        color = tokens["paper-link"]
        return f'<a href="{href}" color="{color}">{content}</a>'
    return content



class SkillTags(Flowable):
    """Selectable text centred by its visible glyph bounds inside rounded tags."""
    def __init__(self, node, tokens):
        super().__init__()
        self.labels = [build.source_text(tag) for tag in node.xpath("./span")]
        self.tokens = tokens
        self.size = measure(tokens, "body-size")
        self.height = measure(tokens, "tag-height")
        self.spaceBefore = 4
        self.widths = [stringWidth(label, "JBMono", self.size) + 2 * measure(tokens, "tag-padding")
                       for label in self.labels]
        self.width = sum(self.widths) + (len(self.labels) - 1) * measure(tokens, "tag-gap")
        face = getFont("JBMono").face
        glyph_data = face.get_table("glyf")
        bounds = [unpack_from(">hhhh", glyph_data, face.glyphPos[face.charToGlyph[ord(char)]] + 2)
                  for label in self.labels for char in label if not char.isspace()]
        scale = self.size / face.unitsPerEm
        bottom = min(bound[1] for bound in bounds) * scale
        top = max(bound[3] for bound in bounds) * scale
        self.baseline = (self.height - (top - bottom)) / 2 - bottom

    def wrap(self, available_width, available_height):
        if self.width > available_width:
            raise ValueError("Certification tags exceed the available line width")
        return self.width, self.height

    def draw(self):
        canvas, x = self.canv, 0
        canvas.saveState()
        canvas.setFont("JBMono", self.size)
        canvas.setLineWidth(0.5)
        for label, width in zip(self.labels, self.widths):
            canvas.setStrokeColor(HexColor(self.tokens["swe-tag-border"]))
            canvas.roundRect(x, 0, width, self.height, measure(self.tokens, "tag-radius"), fill=0, stroke=1)
            canvas.setFillColor(HexColor(self.tokens["swe-tag-text"]))
            canvas.drawCentredString(x + width / 2, self.baseline, label)
            x += width + measure(self.tokens, "tag-gap")
        canvas.restoreState()


def build_pdf(source, output_path, tokens):
    output_path.parent.mkdir(parents=True, exist_ok=True)
    build.register_fonts()
    styles = build.pdf_styles(tokens, source.variant)
    styles["title"].spaceAfter = measure(tokens, "title-gap")
    styles["title"].fontName = "JBMono-Bold"
    styles["section"].fontName = "JBMono-Bold"
    styles["contact"].leading = measure(tokens, "contact-leading")
    styles["job_heading"].textColor = styles["body"].textColor
    styles["bullet"].spaceBefore = measure(tokens, "bullet-gap")
    styles["bullet"].leftIndent = measure(tokens, "bullet-indent")
    styles["bullet"].firstLineIndent = -measure(tokens, "bullet-indent")
    styles["compact"].spaceBefore = measure(tokens, "bullet-gap")
    styles["first-bullet"] = ParagraphStyle("FirstBullet", parent=styles["bullet"], spaceBefore=0)
    styles["sub-bullet"] = ParagraphStyle(
        "SubBullet", parent=styles["bullet"], leftIndent=measure(tokens, "sub-bullet-indent"),
    )
    
    
    def paragraph(node, style):
        return Paragraph(markup(node, tokens), styles[style])
    
    
    story = [
        paragraph(source.article.xpath(".//h1")[0], "title"),
    ]
    for index, contact in enumerate(build.select(source.article, "p.contact")):
        if index:
            story.append(Spacer(1, measure(tokens, "contact-gap")))
        story.append(paragraph(contact, "contact"))
    story.extend([
        Spacer(1, build.token_float(tokens, "pdf-v2-header-summary-gap")),
        paragraph(build.select(source.article, "p.summary")[0], "summary"),
    ])
    
    for section in source.article.xpath("./section"):
        heading = section.xpath("./h2")[0]
        story.extend([
            Spacer(1, build.token_float(tokens, "pdf-v2-section-gap")),
            Paragraph(escape(build.source_text(heading).upper()), styles["section"]),
            Spacer(1, build.token_float(tokens, "pdf-v2-section-content-gap")),
        ])
        jobs = section.xpath("./article[@class='job']")
        if not jobs:
            for index, node in enumerate(section.xpath("./p")):
                style = "body" if index == 0 or "skills-copy" in node.get("class", "").split() else "compact"
                story.append(SkillTags(node, tokens) if "certification-tags" in node.get("class", "").split()
                             else paragraph(node, style))
            continue
    
        for job_index, job in enumerate(jobs):
            if job_index:
                story.append(Spacer(1, build.token_float(tokens, "pdf-v2-job-gap")))
            bullets = build.select(job, "p.bullet")
            # Keep each heading with its first bullet, allowing the rest of a role
            # to continue on the next page instead of leaving a large empty area.
            story.append(KeepTogether([
                paragraph(build.select(job, "p.job-heading")[0], "job_heading"),
                Spacer(1, build.token_float(tokens, "pdf-v2-role-content-gap")),
                paragraph(bullets[0], "first-bullet"),
            ]))
            index = 1
            while index < len(bullets):
                node = bullets[index]
                classes = node.get("class", "").split()
                if "subheading" in classes:
                    story.append(KeepTogether([
                        paragraph(node, "bullet"),
                        paragraph(bullets[index + 1], "sub-bullet"),
                    ]))
                    index += 2
                else:
                    style = "sub-bullet" if "sub-bullet" in classes else "bullet"
                    story.append(KeepTogether([paragraph(node, style)]))
                    index += 1
    
    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=build.A4,
        # ReportLab frames add 6 pt padding inside each document margin.
        leftMargin=measure(tokens, "horizontal-margin") - 6,
        rightMargin=measure(tokens, "horizontal-margin") - 6,
        topMargin=measure(tokens, "vertical-margin") - 6,
        bottomMargin=measure(tokens, "vertical-margin") - 6,
        title=source.title,
        author=source.author,
        subject=source.subject,
        invariant=1,
        pageCompression=1,
    )
    doc.build(story, canvasmaker=partial(
        build.ResumeCanvas, title=source.title, author=source.author, subject=source.subject,
    ))
